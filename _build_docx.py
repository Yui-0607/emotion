from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ── Page setup ──
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# ── Style setup (google_docs_default: Arial-based clean) ──
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15
style.element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
style.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

for level in [1, 2, 3]:
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Arial'
    hs.element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    hs.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

hs1 = doc.styles['Heading 1']
hs1.font.size = Pt(20)
hs1.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
hs1.font.bold = True
hs1.paragraph_format.space_before = Pt(18)
hs1.paragraph_format.space_after = Pt(6)

hs2 = doc.styles['Heading 2']
hs2.font.size = Pt(14)
hs2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
hs2.font.bold = True
hs2.paragraph_format.space_before = Pt(14)
hs2.paragraph_format.space_after = Pt(4)

hs3 = doc.styles['Heading 3']
hs3.font.size = Pt(12)
hs3.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
hs3.font.bold = True
hs3.paragraph_format.space_before = Pt(10)
hs3.paragraph_format.space_after = Pt(2)

def add_para(text, bold=False, style_name='Normal', size=None, color=None, align=None):
    p = doc.add_paragraph(style=style_name)
    run = p.add_run(text)
    if bold: run.bold = True
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor(*color)
    if align is not None: p.alignment = align
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    return p

# ── Title ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
run = title.add_run('见微 · 人生观察者')
run.font.name = 'Arial'
run.font.size = Pt(26)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
run.bold = False
run.element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
run.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run('产品需求文档 (PRD)')
run.font.name = 'Arial'
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
run.element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
run.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')

# ── Section 1 ──
doc.add_heading('1. 产品定位', level=1)
add_para('一个本地运行的个人情绪追踪 + AI 自我观察工具。用户每日记录心情，AI 顺着时间线轻声说出情绪的变化，帮用户看见自己没注意到的模式。')

# ── Section 2 ──
doc.add_heading('2. 技术方案', level=1)
add_bullet('单文件 HTML，数据存 localStorage')
add_bullet('AI 调用 DeepSeek API（deepseek-chat 模型）')
add_bullet('可选 PowerShell 本地服务器（也可直接双击 HTML 使用）')

# ── Section 3 ──
doc.add_heading('3. 功能模块', level=1)

doc.add_heading('3.1 每日记录', level=2)
add_bullet('状态评分：1-10 滑块，两端标签「沉郁 ↔ 澄澈」')
add_bullet('情绪标签：快乐 / 焦虑 / 孤独 / 满足 / 迷茫 / 愤怒，三列网格单选')
add_bullet('自定义情绪：20 种 emoji 选择器 + 自由文本')
add_bullet('日记正文：多行输入，上限 2000 字')

doc.add_heading('3.2 记忆追问', level=2)
add_para('触发时机：保存日记后自动执行', bold=True)
add_bullet('数据范围：最近 5 条记录')
add_bullet('数据格式：按时间正序排列，标注相对日期（今天 / 昨天 / 前天…）')
add_bullet('AI 行为：从前往后读，感受情绪如何一天天流动，说出变化的过程')
add_bullet('口吻：不提问、不说教、不用问号，像朋友轻声的观察')
add_bullet('语言：朴实直白，不用修饰词和比喻，简单有力量')
add_bullet('容错：网络失败自动重试 3 次（间隔 2s / 4s）')

doc.add_heading('3.3 手动观察', level=2)
add_para('触发时机：用户点击「做一次观察」', bold=True)
add_bullet('数据范围：全部历史记录（≥5 条才可触发）')
add_bullet('AI 行为：生成 2-3 段深度观察，关注反复出现的事、情绪转折、被忽视的成长')
add_bullet('口吻与语言：同记忆追问')

doc.add_heading('3.4 定期自动回望', level=2)
add_para('触发时机：每次打开页面自动检查', bold=True)

# Table for periodic observations
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.LEFT
headers = ['类型', '时间跨度', '最少记录']
data = [
    ['本周回望', '≥7 天', '3 条'],
    ['月度回望', '≥30 天', '7 条'],
    ['半年回望', '≥180 天', '15 条'],
    ['年度回望', '≥365 天', '30 条'],
]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = ''
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)

for r, row_data in enumerate(data):
    for c, val in enumerate(row_data):
        cell = table.rows[r+1].cells[c]
        cell.text = ''
        run = cell.paragraphs[0].add_run(val)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)

add_para('')  # spacer
add_bullet('每条期限只生成一次，有新记录自动刷新')

doc.add_heading('3.5 隐藏功能', level=2)
add_bullet('每条记录可隐藏 / 显示')
add_bullet('隐藏后卡片半透明保留')
add_bullet('隐藏的记录不被发送给 AI（追问、回望、观察均自动过滤）')

# ── Section 4 ──
doc.add_heading('4. 数据流', level=1)
add_para('用户保存日记 → localStorage 存储 → 触发 askFU → 取最近 5 条 → 调 DeepSeek → 显示观察')
add_para('用户打开页面 → checkPeriodic → 判断期限 → 调 DeepSeek → 显示回望')
add_para('用户点击观察 → genObs → 取全部记录 → 调 DeepSeek → 显示深度观察')

# ── Section 5 ──
doc.add_heading('5. 已修复的 Bug', level=1)
bugs = [
    ['Emoji 乱码', '.split("") 拆散 surrogate pairs', '改用 [...] 展开运算符'],
    ['追问崩溃', '重试函数缺 async 关键字', '补充 async function rf'],
    ['start.bat 假死', '脚本阻塞等待服务器', '后台启动 + 延迟开浏览器'],
    ['API Key 未转发', '代理服务器空代码块', '正确读写 Authorization header'],
]
bug_table = doc.add_table(rows=5, cols=3)
bug_table.style = 'Table Grid'
bug_headers = ['问题', '根因', '修复']
for i, h in enumerate(bug_headers):
    cell = bug_table.rows[0].cells[i]
    cell.text = ''
    run = cell.paragraphs[0].add_run(h)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)

for r, row_data in enumerate(bugs):
    for c, val in enumerate(row_data):
        cell = bug_table.rows[r+1].cells[c]
        cell.text = ''
        run = cell.paragraphs[0].add_run(val)
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)

# ── Save ──
outpath = r"C:\Users\蒋有成\Documents\emotion\PRD_见微_人生观察者.docx"
doc.save(outpath)
print(f"Saved: {outpath}")
